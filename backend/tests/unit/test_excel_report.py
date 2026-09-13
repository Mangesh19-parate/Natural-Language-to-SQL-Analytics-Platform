import os
import pytest
import openpyxl
from app.services.report_generator import ReportGeneratorService
from app.schemas.report import ReportExportRequest, ReportQueryItem


@pytest.fixture
def report_service(tmp_path):
    return ReportGeneratorService(reports_dir=str(tmp_path))


def test_excel_report_single_query_structure(report_service):
    """
    Test Excel generation for a single query (REQ-RPT-01 / Rule R8.3).
    Must create formatted Summary sheet and dedicated Data sheet.
    """
    query_item = ReportQueryItem(
        query_id="q-002",
        question="Find customer orders above $500",
        executed_sql="SELECT order_id, customer_id, total_amount FROM orders WHERE total_amount > 500;",
        status="success",
        latency_ms=12,
        row_count=2,
        reliability_breakdown={
            "composite_score": 0.94,
            "sub_scores": {
                "schema_grounding": 1.0,
                "join_confidence": 1.0,
                "filter_interpretation": 1.0,
                "execution_validation": 1.0,
                "result_sanity": 0.7,
            },
        },
        columns=["order_id", "customer_id", "total_amount"],
        rows=[
            {"order_id": 101, "customer_id": 12, "total_amount": 750.50},
            {"order_id": 102, "customer_id": 19, "total_amount": 1200.00},
        ],
    )

    req = ReportExportRequest(
        title="High-Value Orders Report",
        scope="single_query",
        role_name="Admin",
        data_source_name="Northwind DB",
        queries=[query_item],
        include_raw_data=True,
    )

    report_id, file_path, content_hash = report_service.generate_excel(req)

    assert os.path.exists(file_path)
    assert file_path.endswith(".xlsx")
    assert len(content_hash) == 64

    # Validate Workbook contents
    wb = openpyxl.load_workbook(file_path)
    assert "Executive Summary" in wb.sheetnames
    assert "Query Results" in wb.sheetnames

    ws_summary = wb["Executive Summary"]
    assert ws_summary["A1"].value == "High-Value Orders Report"

    ws_data = wb["Query Results"]
    assert ws_data.cell(row=4, column=1).value == "order_id"
    assert ws_data.cell(row=4, column=2).value == "customer_id"
    assert ws_data.cell(row=4, column=3).value == "total_amount"
    assert ws_data.cell(row=5, column=1).value == 101


def test_excel_report_multi_query_session(report_service):
    """
    Test session-scope Excel export creating dedicated sheets per query.
    """
    q1 = ReportQueryItem(
        question="Query 1",
        executed_sql="SELECT 1 AS num;",
        columns=["num"],
        rows=[{"num": 1}],
    )
    q2 = ReportQueryItem(
        question="Query 2",
        executed_sql="SELECT 2 AS num;",
        columns=["num"],
        rows=[{"num": 2}],
    )

    req = ReportExportRequest(
        title="Multi-Query Session Workbook",
        scope="session",
        queries=[q1, q2],
    )

    report_id, file_path, content_hash = report_service.generate_excel(req)
    wb = openpyxl.load_workbook(file_path)
    assert "Executive Summary" in wb.sheetnames
    assert "Query_1_Data" in wb.sheetnames
    assert "Query_2_Data" in wb.sheetnames
